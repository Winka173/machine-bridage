"""Prompt 25 H.3 (task F1, DECISIONS 25E): the balance spreadsheet as applied.

    python Tools/balance/export_applied_xlsx.py

Writes Docs/balance/Machine_Brigade_Can_bang_applied.xlsx: a copy of the owner's Docs/balance/Machine_Brigade_Can_bang.xlsx
with a "Hiện tại" column (green) next to each proposed value, holding the game's number after prompt 25, so the owner
can set the two side by side. The numbers come from the game data only: balance.json (weapons with their families and
inherits, vehicles, bosses as the loader builds them, supports, the damage table), campaign.json (unlocks),
Progression.cs (starters, early prices, story loot), the Hud text tables (names) and Docs/backlog/new_content.json
(new content); the two text sheets ("Thay đổi chi tiết", "Kiểm tra từng mục") also get a "Kết quả áp" column with each
row's outcome from Docs/balance/apply-report.md. Nothing runs the game.

The copy holds values: every formula of the owner's file is replaced by the value it computed (openpyxl data_only), since
columns are inserted and a formula's references would no longer point where they did. Styles, widths, frozen panes and
filters are kept. A first sheet, "Ghi chú áp dụng", says what the columns are. A second run writes the same values.
"""
from __future__ import annotations

import copy
import math
import os
import re
import sys
from datetime import date

import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter, range_boundaries

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
sys.path.insert(0, os.path.join(HERE, "..", "docs"))
import import_unlocks as U  # noqa: E402
import import_xlsx as X  # noqa: E402
import report_summary as R  # noqa: E402
import steps_c as C  # noqa: E402
from jsonc_edit import Doc  # noqa: E402

ROOT = X.ROOT
SRC = X.XLSX
OUT = os.path.join(ROOT, "Docs", "balance", "Machine_Brigade_Can_bang_applied.xlsx")
BACKLOG = os.path.join(ROOT, "Docs", "backlog", "new_content.json")
HEAD = "Hiện tại"
OUTCOME_HEAD = "Kết quả áp"
FILL = PatternFill("solid", fgColor="E2F0D9")
HEAD_FILL = PatternFill("solid", fgColor="A9D18E")
DASH = "—"
OUTCOME_VI = {"applied": "đã áp", "already": "đã có sẵn", "deferred": "chờ", "skipped": "bỏ qua"}


def r3(x):
    """A number as the sheet keeps it: an int when whole, else rounded to 4 significant places."""
    if x is None:
        return None
    if isinstance(x, float):
        if not math.isfinite(x):
            return None
        if abs(x - round(x)) < 1e-9:
            return int(round(x))
        return float(f"{x:.4g}")
    return x


# ----------------------------------------------------------------------------------------------------------------
# The game after prompt 25, read from its data


class Now:
    def __init__(self):
        self.g = X.Game(Doc(X.BALANCE))
        self.data = self.g.data
        self.bosses = C.expand(self.data)
        self.dt = self.data["damageTable"]
        self.texts = U.texts()
        self.shop = U.shop_rules()
        progression = open(U.PROGRESSION, encoding="utf-8").read()
        self.loot = set(U.story_loot()) | set(U.cs_list(progression, "SheetLoot"))
        self.unlocks = {}
        for m in U.campaign()["missions"]:
            for u in m.get("unlocks") or []:
                self.unlocks.setdefault(u, m)
        text = R.read(R.REPORT).replace("\r\n", "\n")
        counts, rows = R.parse(text)
        self.final = R.final_outcomes(counts, rows, text)

    # -------------------------------------------------------------- weapons

    def weapon(self, wid):
        return self.g.weapon(wid) if wid in self.g.raw_weapons else None

    @staticmethod
    def cadence(w):
        """(mode, rounds a second, rounds a magazine or burst, rest, sustained DPS) in the sheet's terms: the inverse of
        A2's mapping (a magazine's change and a burst's cooldown are the rest plus one gap; a single shot's rest is its
        cooldown, as the sheet shows it), and the sheet's formula for the DPS."""
        mode, rate, n, dps = X.cadence(w)
        cd = w.get("cooldown", 1.0)
        if w.get("clip", 0) > 0:
            rest = w.get("clipReload", 0.0) - cd
        elif w.get("burst", 1) > 1:
            rest = cd - w.get("burstInterval", 0.1)
        else:
            rest = cd
        return mode, rate, n, rest, dps

    def effect(self, w, column):
        """Matchup.Effect: a weapon's multiplier against armour level 0-5, 'air' or 'structure' (penetration step x damage
        type; a round that strikes the roof never overmatches; thermobaric high explosive on structures)."""
        targets = w.get("targets", "Ground")
        dtype = w.get("damageType", "Kinetic")
        indirect = (w.get("minRange", 0) or 0) > 0 or w.get("projectile") in ("Bomb", "Drone")
        roof = bool(w.get("topAttack")) or indirect
        pen = w.get("pen", 0)
        if column == "air":
            if targets not in ("Air", "All"):
                return 0.0
            return self.penetration(pen, 0, False) * self.dt[dtype]["Air"]
        if targets == "Air":
            return 0.0
        if column == "structure":
            t = self.dt[dtype]["Structure"]
            if w.get("thermobaric") and dtype == "HighExplosive":
                t = max(self.dt.get("thermobaric", 2.0), t)
            return self.penetration(pen, 2, not roof) * t
        return self.penetration(pen, int(column), not roof) * self.dt[dtype]["Ground"]

    def penetration(self, pen, armour, overmatch=True):
        steps = self.dt["penetration"]
        last = len(steps) - 1
        step = 2 - (pen - armour)
        if not overmatch:
            step = max(1, step)
        if step <= 0:
            return steps[0]
        if step >= last:
            return steps[last]
        low = int(math.floor(step))
        return steps[low] + (steps[min(last, low + 1)] - steps[low]) * (step - low)

    def dps_vs(self, w, column):
        return self.cadence(w)[4] * self.effect(w, column)

    # -------------------------------------------------------------- units

    def is_boss(self, vid):
        return vid in self.bosses

    def unit(self, vid):
        if vid in self.bosses:
            return self.bosses[vid]
        return self.g.vehicle(vid) if vid in self.g.vehicles_raw else None

    def mounts(self, vid):
        u = self.unit(vid)
        if u is None:
            return []
        return [m for m in C.mounts_of(u) if m and m in self.g.raw_weapons]

    def hp_shown(self, vid):
        """Health as the sheet shows it: after the vehicles' toughness (x2.2), or a boss's (0.85, and a mini's 0.55)."""
        u = self.unit(vid)
        if u is None or "hp" not in u:
            return None
        if vid in self.bosses:
            share = self.data["bossRanks"]["mini"].get("hp", 1.0) if u.get("rank") == "mini" else 1.0
            return u["hp"] * self.data["toughness"]["bosses"] * share
        return u["hp"] * self.data["toughness"]["vehicles"]

    def unit_dps(self, vid, column):
        """A unit's summed DPS in the sheet's terms: every mount; a boss's ordinary weapons x its weaponDamage on the ground
        (C1), what the boss system lays at its own numbers, its crusher left out."""
        u = self.unit(vid) or {}
        k = u.get("weaponDamage", 1.0) if vid in self.bosses and column != "air" else 1.0
        total = 0.0
        for wid in self.mounts(vid):
            w = self.weapon(wid)
            if w.get("melee") and w.get("laid"):
                continue
            total += self.dps_vs(w, column) * (1.0 if w.get("laid") else k)
        return total

    def name(self, vid, lang="vi", short=False):
        t = self.texts.get(("short." if short else "unit.") + vid)
        return (t[1] if lang == "vi" else t[0]) if t else None

    def armour(self, vid):
        u = self.unit(vid) or {}
        faces = X.armour_faces(u)
        return "/".join(str(x) for x in faces)

    def model_size(self, vid):
        u = self.unit(vid) or {}
        if u.get("modelSize"):
            return " × ".join(f"{x:g}" for x in u["modelSize"])
        if vid in self.bosses and u.get("size"):
            return f"size ×{u['size']:g}"
        return None

    def unlock(self, vid):
        if vid in self.shop["starters"]:
            return "Có sẵn"
        if vid in self.shop["premium"]:
            return f"Cao cấp: {self.shop['premium'][vid]:,} xu".replace(",", ".")
        m = self.unlocks.get(vid)
        if m is None:
            return "Không mở"
        c = m.get("chapter", 0)
        where = f"Xen kẽ {'I' * (c - U.INTERLUDE)}" if c > U.INTERLUDE else f"Chương {c}"
        return where + f" ({m['id']})" + ("; chiến lợi phẩm cốt truyện" if vid in self.loot else "")

    def early(self, vid):
        if vid in self.loot or vid in self.shop["starters"]:
            return DASH
        if vid in self.shop["premium"]:
            return self.shop["premium"][vid]
        return self.shop["early"].get(vid, DASH)

    def super_weapon(self, vid):
        u = self.unit(vid) or {}
        aid = u.get("bigAttack")
        if not aid:
            return "không có" + (" (mini boss)" if u.get("rank") == "mini" else "")
        a = next((x for x in self.data.get("bigAttacks", []) if x["id"] == aid), None)
        if a is None:
            return aid
        import prompt25  # Tools/docs: the design document's words for a strike
        late = a.get("late")
        return (f"{aid}: " + "; ".join(prompt25.strike_words(s) for s in a.get("strikes", []))
                + f"; hồi {a.get('cooldown', 0):g} s, cảnh báo {a.get('warn', 0):g} s"
                + (f"; từ pha {late['phase'] + 1}: {late['count']} phát, hồi {late['cooldown']:g} s" if late else ""))

    def support(self, sid):
        s = next((x for x in self.data["supports"] if x["id"] == sid), None)
        if s is None:
            return None
        k = self.data["firepower"]["strikes"]
        bits = [f"{s.get('cp', 0)} CP"]
        if s.get("damage") and s.get("kind") not in ("Repair", "Smoke", "Scan", "Tower", "Reinforce"):
            bits.append((f"{s['count']} × " if s.get("count") else "") + f"{s['damage'] * k:g}")
        r = s.get("blast") or s.get("radius")
        if r:
            bits.append(f"bán kính {r:g}")
        return " · ".join(bits)

    def item(self, vid, item):
        """The game's value for a row of "Kiểm tra từng mục" or "Thay đổi chi tiết", by its item."""
        it = (item or "").strip()
        u = self.unit(vid)
        if u is None:
            return None
        if it.startswith("Tên"):
            return self.name(vid)
        if it.startswith("Máu"):
            hp = self.hp_shown(vid)
            return r3(round(hp)) if hp is not None else None
        if it.startswith("Giáp"):
            return self.armour(vid)
        if it.startswith("Tốc độ xoay"):
            a, b = u.get("turnRate"), u.get("turretTurnRate")
            if a is None:
                return None
            rad = lambda d: d * math.pi / 180
            return f"{a:g} / {b:g} °/s ({rad(a):.1f} / {rad(b):.1f} rad/s)" if b is not None else f"{a:g} °/s"
        if it.startswith("Tốc độ"):
            return u.get("speed")
        if it.startswith("Tầm nhìn"):
            return u.get("vision")
        if it.startswith("Tầm bắn"):
            return r3(max((self.weapon(w).get("range", 0) for w in self.mounts(vid)), default=None))
        if it.startswith("Giá"):
            return u.get("cp")
        if it.startswith("Kích thước"):
            return self.model_size(vid)
        if it.startswith("Nổ khi bị phá"):
            d = u.get("deathExplosion") or {}
            k = self.data["firepower"]["vehicleBlasts"]   # the sheet shows the blast after the vehicles' firepower
            return f"{d.get('radius', 0):g} m · {d.get('damage', 0) * k:g} sát thương" if d else None
        if it.startswith("Ô"):
            f = u.get("fort") or {}
            return f.get("size") or f.get("kind")
        if it.startswith("Vũ khí: "):
            token = it[len("Vũ khí: "):].split()[0].lower()
            for wid in self.mounts(vid):
                w = self.weapon(wid)
                if token and (token in str(w.get("real", "")).lower() or token in wid.lower()):
                    return f"{wid}: " + self.describe(w)
        if it.startswith("Vũ khí") or it == "Dữ liệu vũ khí":
            ws = self.mounts(vid)
            text = ", ".join(ws)
            if "DPS" in it:
                text += f" · DPS vs giáp 3: {self.unit_dps(vid, 3):.0f}"
            return text
        return None

    def describe(self, w):
        """A weapon in the sheet's words: damage a round, rate, magazine or burst, rest, reach, sustained DPS."""
        mode, rate, n, rest, dps = self.cadence(w)
        return (f"{w.get('damage', 0):g}/phát · {rate:.3g} viên/s · {mode}" + (f" {n}" if n > 1 else "")
                + f" · nạp/nghỉ {rest:.3g} s · tầm {w.get('range', 0):g} m · DPS duy trì {dps:.0f}")

    def outcome(self, sheet, vid, item):
        v = self.final.get((sheet, str(vid), str(item).strip()))
        if not v:
            return None
        step, out, detail = v
        return f"{OUTCOME_VI.get(out, out)} ({step}): {detail}"


# ----------------------------------------------------------------------------------------------------------------
# What each sheet gets: {proposal header (a prefix): fn(row) -> value}; extra: a second column after one of them


def weapons_sheet(now):
    def w(row):
        return now.weapon(row.get("id"))

    def cad(i):
        return lambda row: r3(now.cadence(w(row))[i]) if w(row) else None

    def vs(column):
        return lambda row: r3(now.dps_vs(w(row), column)) if w(row) else None
    cols = {
        "Xuyên đề xuất": lambda row: w(row).get("pen") if w(row) else None,
        "Sát thương/phát đề xuất": lambda row: r3(w(row).get("damage")) if w(row) else None,
        "Chế độ đề xuất": cad(0), "Viên/s đề xuất": cad(1), "Viên mỗi băng/loạt đề xuất": cad(2),
        "Nạp/nghỉ đề xuất": cad(3),
        "Tầm đề xuất": lambda row: r3(w(row).get("range")) if w(row) else None,
        "Tốc độ đạn đề xuất": lambda row: r3(w(row).get("projectileSpeed")) if w(row) else None,
        "Nổ lan đề xuất": lambda row: r3(w(row).get("splash", 0)) if w(row) else None,
        "DPS duy trì đề xuất": cad(4),
        "DPS đề xuất vs giáp 0": vs(0),
    }
    for i in range(1, 6):
        cols[f"=giáp {i}"] = vs(i)
    cols["=trên không"] = vs("air")
    cols["=công trình"] = vs("structure")
    return cols


def unit_weapons_sheet(now):
    def w(row):
        return now.weapon(row.get("Vũ khí id"))

    def mounted(row):
        return row.get("Vũ khí id") in now.mounts(row.get("Đơn vị id"))

    def vs(column):
        return lambda row: r3(now.dps_vs(w(row), column)) if w(row) and mounted(row) else (0 if w(row) else None)
    return {
        "Giữ trong đề xuất": lambda row: 1 if mounted(row) else 0,
        "DPS duy trì đề xuất": lambda row: r3(now.cadence(w(row))[4]) if w(row) and mounted(row) else (0 if w(row) else None),
        "vs giáp 1": vs(1), "vs giáp 3": vs(3), "=trên không": vs("air"), "=công trình": vs("structure"),
    }


def unit_dps_sheet(now):
    def vs(column):
        return lambda row: r3(now.unit_dps(row["id"], column)) if now.unit(row.get("id")) else None

    def per_cp(row):
        u = now.unit(row.get("id"))
        return r3(now.unit_dps(row["id"], 3) / u["cp"]) if u and u.get("cp") else None
    return {"DPS đề xuất vs giáp 1": vs(1), "=vs giáp 3": vs(3), "=trên không": vs("air"), "=công trình": vs("structure"),
            "DPS đề xuất vs giáp 3 / CP": per_cp}


def missile_sheet(now):
    def w(row):
        return now.weapon(row.get("id"))
    return {
        "Tốc độ đề xuất": lambda row: r3(w(row).get("projectileSpeed")) if w(row) else None,
        "Bay tới tầm tối đa mới": lambda row: r3(w(row)["range"] / w(row)["projectileSpeed"]) if w(row) and w(row).get("projectileSpeed") else None,
    }


def round_sheet(now):
    def length(row):
        w = now.weapon(row.get("Vũ khí"))
        if not w:
            return None
        if w.get("roundLength"):
            return r3(w["roundLength"])
        o = now.final.get(("Kích thước đạn", str(row.get("Vũ khí")), str(row.get("Tên thật"))))
        m = re.search(r"(?:kept: |-> )(\d+(?:\.\d+)?) m", o[2]) if o else None
        return float(m.group(1)) if m else None
    return {"Dài đề xuất": length}


def cp_sheet(now, key="id"):
    return lambda row: (now.unit(row.get(key)) or {}).get("cp")


def boss_sheet(now):
    types = {k: v["Ground"] for k, v in now.dt.items() if isinstance(v, dict)}
    pens = {2 - i: p for i, p in enumerate(now.dt["penetration"])}   # pen - armour -> its step (DamageTable)

    def dps(row):
        vid = row.get("id")
        if vid not in now.bosses:
            return None
        scaled, fixed, _ = C.boss_dps(now.g, now.bosses[vid], types, pens)
        return r3(round(scaled * now.bosses[vid].get("weaponDamage", 1.0) + fixed))
    return {
        "Máu đề xuất": lambda row: r3(round(now.hp_shown(row["id"]))) if now.hp_shown(row.get("id")) else None,
        "DPS thường mục tiêu": dps,
        "Vũ khí thêm / đổi": lambda row: ", ".join(now.mounts(row.get("id"))) or None,
        "Siêu vũ khí đề xuất": lambda row: now.super_weapon(row["id"]) if now.unit(row.get("id")) else None,
    }


def names_sheet(now):
    return {
        "Tên đầy đủ đề xuất": lambda row: now.name(row.get("id")),
        "Tên ngắn": lambda row: now.name(row.get("id"), short=True),
        "Tên tiếng Anh": lambda row: now.name(row.get("id"), "en"),
    }


def vehicles_sheet(now):
    return {"Mở khóa đề xuất": lambda row: now.unlock(row.get("id")), "Mua sớm": lambda row: now.early(row.get("id"))}


SUPPORT_CARDS = {"Không kích": ["airstrike"], "Pháo kích": ["artillery_barrage"], "Tên lửa hành trình": ["cruise_missile"],
                 "Bom napalm": ["napalm_strike"], "Đòn SEAD": ["sead_strike"],
                 "Màn khói, UAV quét, Sửa chữa, Tháp dã chiến": ["smoke_screen", "uav_scan", "repair_drop", "field_tower"]}


def support_sheet(now):
    return {"Đề xuất": lambda row: "; ".join(filter(None, (now.support(s) for s in SUPPORT_CARDS.get(row.get("Thẻ"), []))))}


def text_sheet(now, sheet, item_col, proposal):
    return {proposal: lambda row: r3(now.item(row.get("id"), row.get(item_col)))}, \
        {proposal: (OUTCOME_HEAD, lambda row: now.outcome(sheet, row.get("id"), row.get(item_col)))}


def round2_sheet(now):
    def out(row):
        o = now.final.get(("Cân bằng lần 2", str(row.get("id")), str(row.get("Kết luận"))))
        return f"{OUTCOME_VI.get(o[1], o[1])} ({o[0]}): {o[2]}" if o else None
    return {"CP sau đợt 1": cp_sheet(now), "Đề xuất đợt 2": out}


def backlog_sheet(now, name_col):
    import json
    items = json.load(open(BACKLOG, encoding="utf-8"))["items"]
    by_name = {}
    for it in items:
        by_name.setdefault(it["name"], it)

    def status(row):
        it = by_name.get(row.get(name_col))
        if not it:
            return None
        return it["status"] + (f" · id {it['id']}" if it.get("id") else "") + (f" · {it['notes']}" if it.get("notes") else "")
    return status


def factors_sheet(now):
    names = {"Động năng": "Kinetic", "Nổ lõm": "ShapedCharge", "Nổ mạnh": "HighExplosive", "Lửa": "Fire",
             "Mảnh": "Fragmentation", "Năng lượng": "Energy"}
    steps = {"≥ +2": 0, "+1": 1, "0": 2, "−1": 3, "−2": 4, "≤ −3": 5}

    def col(kind):
        def f(row):
            label = str(row.get("Loại sát thương") or "").strip()
            if label in names:
                return now.dt[names[label]][kind]
            if kind == "Ground" and label in steps:
                return now.dt["penetration"][steps[label]]
            if kind == "Ground" and label == "Chênh xuyên − giáp":
                return HEAD
            return None
        return f
    return {"Lên mặt đất": col("Ground"), "Lên máy bay": col("Air"), "Lên công trình": col("Structure")}


def plan(now):
    """{sheet: ({header prefix: fn}, {header prefix: (second header, fn)})}."""
    tdct, tdct2 = text_sheet(now, "Thay đổi chi tiết", "Hạng mục", "Đề xuất (cụ thể)")
    ktm, ktm2 = text_sheet(now, "Kiểm tra từng mục", "Mục", "Đề xuất")
    p = {
        "Vũ khí đề xuất": (weapons_sheet(now), {}),
        "Đơn vị – vũ khí": (unit_weapons_sheet(now), {}),
        "Tổng DPS đơn vị": (unit_dps_sheet(now), {}),
        "Tốc độ tên lửa": (missile_sheet(now), {}),
        "Kích thước đạn": (round_sheet(now), {}),
        "Giá CP": ({"CP đề xuất": cp_sheet(now)}, {}),
        "Kích thước – giá": ({"CP sau đề xuất": cp_sheet(now)}, {}),
        "Boss đề xuất": (boss_sheet(now), {}),
        "Boss": ({"Máu đề xuất": boss_sheet(now)["Máu đề xuất"]}, {}),
        "Tên đề xuất": (names_sheet(now), {}),
        "Phương tiện": (vehicles_sheet(now), {}),
        "Thẻ hỗ trợ": (support_sheet(now), {}),
        "Thay đổi chi tiết": (tdct, tdct2),
        "Kiểm tra từng mục": (ktm, ktm2),
        "Cân bằng lần 2": (round2_sheet(now), {}),
        "Hệ số": (factors_sheet(now), {}),
    }
    for sheet, col in (("Thẻ hỗ trợ mới", "Thẻ"), ("Đề xuất thêm", "Tên đề xuất"), ("Công trình mới", "Tên"),
                       ("Tên lửa & bom mới", "Tên"), ("Boss mới", "Tên · dòng phụ")):
        p[sheet] = ({"Mở khóa đề xuất" if sheet != "Boss mới" else "Ra mắt đề xuất": backlog_sheet(now, col)}, {})
    return p


# ----------------------------------------------------------------------------------------------------------------
# Writing


def match(headers, prefix):
    """The columns a plan key means: "=name" is an exact header, else the first header starting with it."""
    if prefix.startswith("="):
        return [i for i, h in enumerate(headers) if h == prefix[1:]]
    return [next(i for i, h in enumerate(headers) if h.startswith(prefix))] if any(h.startswith(prefix) for h in headers) else []


def rebuild(wb, name, inserts):
    """The sheet with new columns: inserts = {old column index (1-based) after which: [(header, [value by row])]}."""
    old = wb[name]
    index = wb.sheetnames.index(name)
    new = wb.create_sheet(name + "~", index)
    after = sorted(inserts)

    def shift(c):
        return c + sum(len(inserts[a]) for a in after if a < c)
    for row in old.iter_rows():
        for cell in row:
            if cell.value is None and not cell.has_style:
                continue
            n = new.cell(row=cell.row, column=shift(cell.column), value=cell.value)
            if cell.has_style:
                n._style = copy.copy(cell._style)
    for a in after:
        for k, (header, values) in enumerate(inserts[a]):
            col = shift(a) + 1 + k
            src_head = old.cell(row=1, column=a)
            h = new.cell(row=1, column=col, value=header)
            if src_head.has_style:
                h._style = copy.copy(src_head._style)
            h.fill = HEAD_FILL
            h.font = Font(bold=True)
            for r, v in values.items():
                c = new.cell(row=r, column=col, value=v)
                src = old.cell(row=r, column=a)
                if src.has_style:
                    c._style = copy.copy(src._style)
                c.fill = FILL
    for key, dim in old.column_dimensions.items():
        lo, hi = dim.min or openpyxl.utils.column_index_from_string(key), dim.max or openpyxl.utils.column_index_from_string(key)
        for c in range(lo, min(hi, old.max_column) + 1):
            if dim.width:
                new.column_dimensions[get_column_letter(shift(c))].width = dim.width
    for a in after:
        width = old.column_dimensions[get_column_letter(a)].width or 14
        for k in range(len(inserts[a])):
            new.column_dimensions[get_column_letter(shift(a) + 1 + k)].width = max(12, min(40, width))
    for r, dim in old.row_dimensions.items():
        if dim.height:
            new.row_dimensions[r].height = dim.height
    if old.freeze_panes:
        c, r = openpyxl.utils.cell.coordinate_from_string(old.freeze_panes)
        new.freeze_panes = f"{get_column_letter(shift(openpyxl.utils.column_index_from_string(c)))}{r}"
    for rng in old.merged_cells.ranges:
        c1, r1, c2, r2 = range_boundaries(str(rng))
        new.merge_cells(start_row=r1, start_column=shift(c1), end_row=r2, end_column=shift(c2))
    if old.auto_filter.ref:
        c1, r1, c2, r2 = range_boundaries(old.auto_filter.ref)
        new.auto_filter.ref = f"{get_column_letter(shift(c1))}{r1}:{get_column_letter(shift(c2) + sum(len(inserts[a]) for a in after if a == c2))}{r2}"
    wb.remove(old)
    new.title = name


def notes_sheet(wb, filled):
    ws = wb.create_sheet("Ghi chú áp dụng", 0)
    lines = [
        ("Bản đã áp (prompt 25)", f"Xuất ngày {date.today().isoformat()} bằng Tools/balance/export_applied_xlsx.py từ dữ liệu game sau khi áp file cân bằng."),
        ("Cột \"Hiện tại\" (nền xanh)", "Số trong game sau khi áp, đặt ngay cạnh cột đề xuất để so. Các cột \"... hiện\" có sẵn trong file là số TRƯỚC khi áp."),
        ("Cột \"Kết quả áp\"", "Sheet Thay đổi chi tiết và Kiểm tra từng mục: kết quả của từng dòng trong Docs/balance/apply-report.md "
                                "(đã áp / đã có sẵn / chờ / bỏ qua, kèm việc và chi tiết)."),
        ("Nguồn số", "balance.json (vũ khí theo họ và kế thừa, xe, boss như game dựng, thẻ hỗ trợ, bảng sát thương), campaign.json (mở khóa), "
                     "Progression.cs (có sẵn, giá mua sớm, chiến lợi phẩm), bảng chữ trong Hud (tên), Docs/backlog/new_content.json (nội dung mới)."),
        ("Cách tính", "Chế độ bắn, viên/s, viên mỗi băng hoặc loạt, nạp/nghỉ: ngược lại quy tắc A2 (thay băng và nghỉ loạt = nghỉ + một nhịp). "
                      "DPS duy trì: công thức của file (sát thương × số viên / (số viên / nhịp + nghỉ)). DPS theo cấp giáp: × bước xuyên × hệ số loại đạn "
                      "(đạn đánh nóc không vượt cấp). Máu: như file hiển thị (xe × 2,2; boss × 0,85, mini boss × 0,55). Tổng DPS boss: vũ khí thường × weaponDamage."),
        ("Luật của game", "Cột \"Hiện tại\" theo luật của game (prompt 15 thắng về hệ thống): vũ khí không ngắm được máy bay hay mặt đất "
                          "thì bằng 0 ở đó; đạn đánh nóc (bắn cầu vồng, thả, đánh nóc) không vượt cấp giáp; nổ mạnh nhiệt áp ×2 lên công trình; "
                          "tổng DPS của xe gồm cả vũ khí mà ghi chú của file thêm vào (không có dòng riêng). Khác với file ở đây là do luật, không phải số."),
        ("Công thức", "Bản này chỉ giữ giá trị: mọi công thức của file gốc được thay bằng giá trị nó đã tính, vì chèn cột làm lệch tham chiếu. File gốc giữ nguyên."),
        ("Số cột thêm", ", ".join(f"{k}: {v}" for k, v in filled.items())),
    ]
    for r, (a, b) in enumerate(lines, 1):
        ws.cell(row=r, column=1, value=a).font = Font(bold=True)
        ws.cell(row=r, column=2, value=b)
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 140


def main():
    now = Now()
    wb = openpyxl.load_workbook(SRC, data_only=True)
    filled = {}
    for name, (cols, extra) in plan(now).items():
        ws = wb[name]
        headers = [str(c.value).strip() if c.value is not None else "" for c in ws[1]]
        rows = {}
        for r in range(2, ws.max_row + 1):
            vals = [ws.cell(row=r, column=c + 1).value for c in range(len(headers))]
            if any(v not in (None, "") for v in vals):
                rows[r] = {h: v for h, v in zip(headers, vals) if h}
        inserts = {}
        for prefix, fn in cols.items():
            for i in match(headers, prefix):
                values = {}
                for r, row in rows.items():
                    try:
                        v = fn(row)
                    except (KeyError, TypeError, ValueError, ZeroDivisionError):
                        v = None
                    if v is not None and v != "":
                        values[r] = v
                inserts.setdefault(i + 1, []).append((HEAD, values))
                if prefix in extra:
                    head2, fn2 = extra[prefix]
                    inserts[i + 1].append((head2, {r: v for r, v in ((r, fn2(row)) for r, row in rows.items()) if v}))
        if not inserts:
            print("nothing matched on", name)
            continue
        rebuild(wb, name, inserts)
        filled[name] = sum(len(v) for v in inserts.values())
    notes_sheet(wb, filled)
    wb.save(OUT)
    print("wrote", OUT, filled)


if __name__ == "__main__":
    main()
