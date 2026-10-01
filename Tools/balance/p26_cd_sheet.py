"""Prompt 26 D.3-D.12: the post-1945 reference lines in the balance spreadsheet and the new-content tracker (DECISIONS 26CD).

Reads Tools/docs/unit_refs.json (the reference lines, already rewritten) and writes, with openpyxl and nothing else touched:

  * the "Tham khảo (ngoài đời / game)" cell of every unit, building and boss row whose line changed;
  * the Boss sheet's "Hình dạng (cho AI vẽ)" (the "Dựa trên: ..." part, and "Kích thước hiện" with the sizes now in the data);
  * Ixion's, Gungnir's and Leviathan's description cells; the new-content sheets' names and reference cells (Flak 88 -> KS-19
    100 mm, the barrage balloon -> the JLENS radar aerostat, Pak 40, Horten);
  * Docs/backlog/new_content.json with the same names (the tracker keeps its statuses).

    python Tools/balance/p26_cd_sheet.py            # dry run (prints what would change)
    python Tools/balance/p26_cd_sheet.py --write

Re-runnable: a second run changes nothing. The sheet's history tabs ("Thay đổi chi tiết", "Kiểm tra từng mục", "Việc cho agent")
are logs of what was asked at the time and are left as written.
"""
import json
import os
import re
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
sys.stdout.reconfigure(encoding="utf-8")

import glb_bounds as G  # noqa: E402
import openpyxl  # noqa: E402
import p26_ab as P  # noqa: E402
from jsonc_edit import Doc  # noqa: E402

ROOT = os.path.abspath(os.path.join(HERE, "..", ".."))
XLSX = os.path.join(ROOT, "Docs", "balance", "Machine_Brigade_Can_bang.xlsx")
REFS = os.path.join(ROOT, "Tools", "docs", "unit_refs.json")
TRACKER = os.path.join(ROOT, "Docs", "backlog", "new_content.json")

REF_IDS = ["drone_mothership", "locust", "flame_tank", "titan_tank", "gun_turret", "heavy_turret", "heavy_turret.coastal", "super_gun", "spawn_bastion", "coastal_battery",
           "behemoth", "behemoth_inferno", "behemoth_tempest", "behemoth_mk2", "mobile_fortress", "fenrir", "rail_supergun", "command_airship",
           "argus", "leviathan", "scylla", "ixion"]

DESC = {
    "ixion": "xe tải mỏ bọc thép BelAZ-75710 chạy thẳng khoảng 7 m/s, xoay rất chậm; tháp pháo 125 mm hàn trên thùng tự đổi đạn nổ mạnh (lõi 5 m, rìa 10 m) và đạn xuyên, khoảng 780 mỗi 2 giây; cứ khoảng 10 giây lao nghiền theo đường báo trước 2 giây, khoảng 900 và choáng 1 giây; hai súng máy 12,7 mm trên nóc; khi rẽ đổ ra sau dải sáu quả mìn tồn tại 20 giây. Giáp trước 4, hông 3, sau và nóc 2; bốn lốp sau giáp 1 là điểm yếu; phá một lốp trước thì cú lao lệch và dừng.",
    "rail_supergun": "cứ 25 giây một phát điện từ vào cụm xe mặt đất đông nhất của bạn, ở bất cứ đâu trên bản đồ: xuyên qua tối đa năm xe trên một đường thẳng rồi nổ ở xe cuối; vòng đỏ báo điểm ngắm trước 3 giây. trừng phạt đội quân dồn cục; không di chuyển được, và mất trạm chỉ thị thì phát bắn rơi lệch xa.",
}


def ref_text(entry):
    text = ", ".join(entry["real"])
    if entry["media"]:
        text += " · " + ", ".join(entry["media"])
    if entry["note"]:
        text += " — " + entry["note"]
    return text


def dims():
    """id -> 'L × W × H' of every boss now in the data (the native model box times the drawn scale)."""
    data = Doc(P.PATH).data()
    built = P.expand(data)
    raw = {v["id"]: v for v in data["vehicles"]}

    def model_of(i):
        return built[i].get("model") or (model_of(raw[i]["variantOf"]) if "variantOf" in raw[i] else i)

    def final(i):
        r = raw[i]
        if "variantOf" in r:
            return final(r["variantOf"]) * (r.get("variant") or {}).get("size", 0.7)
        return built[i].get("scale", 1.0) * r.get("size", 1.0)

    out = {}
    for i, b in built.items():
        if not b.get("boss"):
            continue
        if raw[i].get("modelSize"):
            L, W, H = raw[i]["modelSize"]
        else:
            w, h, l = G.bounds(model_of(i))
            s = final(i)
            L, W, H = l * s, w * s, h * s
        out[i] = f"{L:.1f} × {W:.1f} × {H:.1f}"
    return out


def main():
    refs = json.load(open(REFS, encoding="utf-8"))
    size = dims()
    wb = openpyxl.load_workbook(XLSX)
    changes = []

    def put(ws, cell, value):
        if cell.value != value:
            changes.append(f"{ws.title}!{cell.coordinate}")
            cell.value = value

    for name in ("Phương tiện", "Công trình", "Boss"):
        ws = wb[name]
        head = {str(c.value).strip(): c.column for c in ws[1] if c.value}
        ref_col = head["Tham khảo (ngoài đời / game)"]
        shape_col = head.get("Hình dạng (cho AI vẽ)")
        desc_col = head.get("Miêu tả")
        for r in range(2, ws.max_row + 1):
            vid = ws.cell(r, 1).value
            if not vid:
                continue
            if vid in REF_IDS and vid in refs:
                put(ws, ws.cell(r, ref_col), ref_text(refs[vid]))
            if name == "Boss":
                shape = ws.cell(r, shape_col)
                if isinstance(shape.value, str):
                    v = shape.value
                    if vid in REF_IDS and vid in refs:
                        v = re.sub(r"^Dựa trên: .*?\. (?=Kích thước hiện)|^Dựa trên: .*?(?=Kích thước hiện)", "Dựa trên: " + ref_text(refs[vid]) + ". ", v, count=1, flags=re.S)
                    if vid in size:
                        v = re.sub(r"Kích thước hiện[^;]*? m(?=;)", f"Kích thước hiện {size[vid]} m", v, count=1)
                    put(ws, shape, v)
                if vid in DESC:
                    for col in (desc_col, head["Cách đánh"]):
                        put(ws, ws.cell(r, col), DESC[vid])
                if vid == "leviathan":
                    for col in (desc_col, head["Cách đánh"], head["Vũ khí"], head["Đề xuất"], head["Lý do thay đổi"]):
                        cell = ws.cell(r, col)
                        if isinstance(cell.value, str):
                            put(ws, cell, cell.value.replace("Type 94 460 mm/45", "Mk 7 406 mm/50 (Iowa,").replace("460 mm", "406 mm").replace("Mk 7 406 mm/50 (Iowa, (triple)", "Mk 7 406 mm/50 (Iowa, triple)"))

    # the new-content sheets
    flak_name = "Tháp cao xạ hạng nặng 100 mm (KS-19)"
    balloon_name = "Khí cầu neo radar (JLENS)"
    for sheet, edits in {
        "Đề xuất thêm": {"A2": flak_name, "G2": "KS-19 100 mm (1947)", "H2": "Red Alert 2 (Flak Cannon)", "H7": "Company of Heroes (pháo chống tăng 57 mm)",
                         "A30": balloon_name, "G30": "JLENS (khí cầu neo radar)"},
        "Công trình mới": {"A2": flak_name, "V2": "KS-19 100 mm (1947)", "W2": "Red Alert 2 (Flak Cannon)", "W3": "Company of Heroes (pháo chống tăng)",
                           "A10": balloon_name, "V10": "JLENS (khí cầu neo radar)"},
        "Boss mới": {"I3": "2B1 Oka và Object 271 (1957)", "I4": "Cánh bay B-2 phóng to; B-21 Raider"},
    }.items():
        ws = wb[sheet]
        for coord, value in edits.items():
            put(ws, ws[coord], value)
    for sheet, coord in (("Đề xuất thêm", "I3"), ("Công trình mới", "X15")):
        cell = wb[sheet][coord]
        if isinstance(cell.value, str):
            put(wb[sheet], cell, cell.value.replace("Flak 88", "KS-19 100 mm"))

    tracker = json.load(open(TRACKER, encoding="utf-8"))
    for it in tracker["items"]:
        if it["key"] == "dx01":
            it["name"] = flak_name
            it["notes"] = (it.get("notes") or "").split(" | 26CD")[0] + " | 26CD: a KS-19 100 mm (1947) heavy AA tower, slow burst, big blast (the mechanism is kept)."
        if it["key"] == "dx29":
            it["name"] = balloon_name
            it["notes"] = (it.get("notes") or "").split(" | 26CD")[0] + " | 26CD: a JLENS radar aerostat; keeps the bomb scatter and reveals stealth aircraft near it."
    print(f"{len(changes)} cells: " + ", ".join(changes))
    if "--write" in sys.argv:
        wb.save(XLSX)
        with open(TRACKER, "w", encoding="utf-8", newline="") as f:
            json.dump(tracker, f, ensure_ascii=False, indent=2)
            f.write("\n")
        print("written")
    else:
        print("dry run (--write saves)")


if __name__ == "__main__":
    main()
