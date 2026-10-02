"""Prompt 28 O.7: the AI research sheet as applied.

    python Tools/ai/export_applied_xlsx.py

Writes Docs/ai/Machine_Brigade_AI_Research_applied.xlsx: a copy of Docs/ai/Machine_Brigade_AI_Research.xlsx (values; its
formulas replaced by what they computed, openpyxl data_only) with the game's numbers beside the sheet's, so the owner
sees what the game plays with: green columns on "Tham số AI" (value, min, max and the balance.json key), on "Cân bằng
kinh tế" (the game's value) and on "Chiến thuật" (id, merged into, preferred card ids), a "Tham số thế giới" sheet (the
World Model's own parameters, not in the research sheet) and a first sheet "Ghi chú áp dụng". Everything comes from
balance.json; nothing runs the game. Re-run it after the sweeps change values or merge tactics.
"""
from __future__ import annotations

import json
import os
import re
import sys
from datetime import date

import openpyxl
from openpyxl.styles import Font, PatternFill

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import import_ai_xlsx as I  # noqa: E402

OUT = os.path.join(I.ROOT, "Docs", "ai", "Machine_Brigade_AI_Research_applied.xlsx")
GREEN = PatternFill("solid", fgColor="C6EFCE")
BOLD = Font(bold=True)


def balance() -> dict:
    """balance.json without its comments (whole-line and trailing ones outside strings)."""
    with open(I.BALANCE, encoding="utf-8") as f:
        text = f.read()
    out = []
    for line in text.split("\n"):
        stripped, in_string, i = [], False, 0
        while i < len(line):
            c = line[i]
            if c == '"' and (i == 0 or line[i - 1] != "\\"):
                in_string = not in_string
            if not in_string and line.startswith("//", i):
                break
            stripped.append(c)
            i += 1
        out.append("".join(stripped))
    text = re.sub(r",(\s*[}\]])", r"\1", "\n".join(out))
    return json.loads(text)


def add_columns(ws, header_row, titles):
    first = ws.max_column + 1
    for k, t in enumerate(titles):
        cell = ws.cell(row=header_row, column=first + k, value=t)
        cell.fill = GREEN
        cell.font = BOLD
    return first


def main():
    data = balance()
    ai, beh = data["ai"], data.get("aiBehaviour", {})
    wb = openpyxl.load_workbook(I.XLSX, data_only=True)

    # Tham số AI: the game's value, min, max and key beside each row.
    ws = wb["Tham số AI"]
    col = add_columns(ws, 1, ["Game: giá trị", "Game: nhỏ nhất", "Game: lớn nhất", "Khóa balance.json"])
    for r in range(2, ws.max_row + 1):
        name = str(ws.cell(row=r, column=1).value or "")
        key = next((k for n, k, _ in I.PARAMS if name.startswith(n)), None)
        if key is None or key not in ai["params"]:
            continue
        p = ai["params"][key]
        for k, v in enumerate([p["value"], p["min"], p["max"], f"ai.params.{key}"]):
            ws.cell(row=r, column=col + k, value=v).fill = GREEN

    # Cân bằng kinh tế: the game's value by the row's label.
    ws = wb["Cân bằng kinh tế"]
    econ = ai["economy"]
    labels = {"Thu nhập CP gốc": "baseIncome", "Thêm mỗi cứ điểm": "perPoint", "Trần quân mặt đất": "groundCap",
              "Trần máy bay": "airCap", "Trần CP tích trữ": "stockCap", "Ngưỡng tấn công của tướng AI": "attackThreshold",
              "Ngưỡng tấn công thấp nhất": "useOrLoseThreshold", "Thời gian giảm ngưỡng": "useOrLoseRamp"}
    bands = {"Ít quân": 0, "Vừa": 1, "Đông": 2}
    col = None
    for r in range(1, ws.max_row + 1):
        label = str(ws.cell(row=r, column=1).value or "").strip()
        if label == "Tham số":
            col = add_columns(ws, r, ["Game"])
        value = None
        for prefix, key in labels.items():
            if label.startswith(prefix):
                value = econ[key]["value"]
        if label.startswith("Mốc bậc 1"):
            value = ", ".join(str(t["value"]) for t in econ["escalation"])
        if value is not None and col is not None:
            ws.cell(row=r, column=col, value=value).fill = GREEN

    # Chiến thuật: id, merge, preferred card ids, modules.
    ws = wb["Chiến thuật"]
    col = add_columns(ws, 1, ["Game: id", "Game: gộp vào", "Game: ưu tiên mua (id)", "Game: mô-đun hành vi"])
    by_name = {t["name"]: t for t in beh.get("tactics", [])}
    for r in range(2, ws.max_row + 1):
        t = by_name.get(str(ws.cell(row=r, column=1).value or "").strip())
        if t is None:
            continue
        values = [t["id"], t.get("mergedInto", ""), ", ".join(t["prefer"]), json.dumps(t["modules"], ensure_ascii=False)]
        for k, v in enumerate(values):
            ws.cell(row=r, column=col + k, value=v).fill = GREEN

    # The World Model's own parameters.
    ws = wb.create_sheet("Tham số thế giới")
    ws.append(["Khóa balance.json", "Giá trị", "Nhỏ nhất", "Lớn nhất", "Chỉ số đo", "Lý do"])
    for c in ws[1]:
        c.font = BOLD
        c.fill = GREEN
    for key, p in ai["world"].items():
        ws.append([f"ai.world.{key}", p["value"], p["min"], p["max"], p["metric"], p["reason"]])

    note = wb.create_sheet("Ghi chú áp dụng", 0)
    for line in [
        f"Bản áp dụng của file nghiên cứu AI (prompt 28 O.7), xuất ngày {date.today().isoformat()} từ balance.json.",
        "Cột màu xanh: giá trị game đang dùng (Tools/ai/export_applied_xlsx.py). Các cột khác là bản gốc (giá trị, không còn công thức).",
        "Các giá trị là khởi điểm từ file gốc; mô phỏng (AiParamSweep, TacticFingerprintSweep) sẽ chỉnh khi được cho phép chạy.",
        "Gộp chiến thuật (H.3): cột 'Game: gộp vào' trên sheet Chiến thuật; trống = chưa gộp.",
        "Sheet 'Tham số thế giới': tham số riêng của World Model và áp lực tăng dần, không có trong file gốc (DECISIONS 28 1, 28 5).",
    ]:
        note.append([line])
    note.column_dimensions["A"].width = 130
    wb.save(OUT)
    print(f"written {OUT}")


if __name__ == "__main__":
    main()
