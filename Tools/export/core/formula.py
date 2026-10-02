"""Layer B (spec 2.3): live Excel formulas, a small evaluator for the subset the domains write, and the two tests:
formula == the value the game's code computes (its Python port, within 1e-6) and input_<name> copy == its source file.

A formula cell is an F object put in a row like any value:
    row.set("chu_ky_day_du_s", F("=IF({Vu_khi!bang_dan}>0,...)", expect=py_value, ref="game: Sim/Content/Definitions.cs CycleSeconds"))
Placeholders (resolved once every book is built, when the column letters and row numbers are final):
    {col}                 this sheet, this row          {Sheet!col}        sheet Sheet of the same file, the row of the same id
    {Sheet!col@<id>}      the row of that id            {Sheet!col@*}      the whole column (rows 2..n+1), for ranges
    {col@<id>} / {col@*}  the same in this sheet
Formulas only reference sheets of the same file (spec 2.2). The xlsx gets the formula; the CSV gets its evaluated value.
"""
from __future__ import annotations

import math
import re

from openpyxl.utils import get_column_letter

TOL = 1e-6
PLACEHOLDER = re.compile(r"\{(?:([A-Za-z][A-Za-z0-9_]*)!)?([A-Za-z0-9_]+)(?:@([^}]*))?\}")


class FormulaError(Exception):
    pass


class F:
    """A live formula cell. expect: the reference value (the game's, ported; or the exporter's own Python for an analysis
    column); ref: where expect comes from ('game: <C# file> <method>' or 'python: <tool>')."""
    __slots__ = ("template", "expect", "ref", "text", "value", "error", "_busy")

    def __init__(self, template: str, expect=None, ref: str = ""):
        self.template = template if template.startswith("=") else "=" + template
        self.expect = expect
        self.ref = ref
        self.text = None
        self.value = None
        self.error = None
        self._busy = False

    def __repr__(self):
        return f"F({self.text or self.template!r})"


def ref(sheet: str, col: str, at: str | None = None) -> str:
    """A placeholder text: {sheet!col} / {sheet!col@at}; sheet '' = this sheet."""
    s = f"{sheet}!" if sheet else ""
    return "{" + s + col + (f"@{at}" if at is not None else "") + "}"


def lookup(sheet: str, col: str, key_expr: str, id_col: str = "id") -> str:
    """INDEX(sheet.col, MATCH(key, sheet.id, 0)): a value of the row whose id is key (a live lookup)."""
    return f"INDEX({ref(sheet, col, '*')},MATCH({key_expr},{ref(sheet, id_col, '*')},0))"


def q(text: str) -> str:
    """An Excel string literal."""
    return '"' + str(text).replace('"', '""') + '"'


# ---------------------------------------------------------------------------------------------------------- resolve
def _layout(books):
    lay = {}
    for fid, book in books.items():
        for sname, s in book.sheets.items():
            cols = s.column_order()
            rows = s.sorted_rows()
            lay[(fid, sname)] = ({c: get_column_letter(i + 1) for i, c in enumerate(cols)},
                                 {str(r.id): k + 2 for k, r in enumerate(rows)}, len(rows), s)
    return lay


def resolve(books) -> list[str]:
    """Turns every F's template into A1 text. Returns the problems (a missing sheet, column or row)."""
    lay = _layout(books)
    problems = []
    for fid, book in books.items():
        for sname, s in book.sheets.items():
            for r in s.rows.values():
                for cname, v in r.values.items():
                    if not isinstance(v, F):
                        continue

                    def sub(m, fid=fid, sname=sname, r=r, cname=cname):
                        sh, col, at = m.group(1), m.group(2), m.group(3)
                        key = (fid, sh or sname)
                        if key not in lay:
                            problems.append(f"{fid}/{sname}.{cname} row {r.id}: no sheet {sh}")
                            return "#REF!"
                        letters, rows, n, _s = lay[key]
                        if col not in letters:
                            problems.append(f"{fid}/{sname}.{cname} row {r.id}: no column {sh or sname}.{col}")
                            return "#REF!"
                        prefix = f"{sh}!" if sh and sh != sname else ""
                        L = letters[col]
                        if at == "*":
                            return f"{prefix}${L}$2:${L}${max(2, n + 1)}"
                        rid = str(r.id) if at is None else at
                        if rid not in rows:
                            problems.append(f"{fid}/{sname}.{cname} row {r.id}: no row {rid} in {sh or sname}")
                            return "#REF!"
                        return f"{prefix}${L}${rows[rid]}"

                    v.text = PLACEHOLDER.sub(sub, v.template)
    return problems


# ---------------------------------------------------------------------------------------------------------- evaluate
TOKEN = re.compile(r"""
    (?P<ws>\s+)
  | (?P<str>"(?:[^"]|"")*")
  | (?P<ref>(?:[A-Za-z_][A-Za-z0-9_]*!)?\$?[A-Z]{1,3}\$?\d+(?::\$?[A-Z]{1,3}\$?\d+)?)
  | (?P<num>\d+(?:\.\d*)?(?:[eE][+-]?\d+)?|\.\d+)
  | (?P<func>[A-Z][A-Z0-9.]*(?=\())
  | (?P<bool>TRUE|FALSE)
  | (?P<op><=|>=|<>|[-+*/^&=<>(),])
""", re.X)
CELL = re.compile(r"\$?([A-Z]{1,3})\$?(\d+)")


def _col_num(letters: str) -> int:
    n = 0
    for ch in letters:
        n = n * 26 + ord(ch) - 64
    return n


class Range(list):
    pass


class _Eval:
    def __init__(self, books):
        self.books = books
        self.grids = {}
        self.ranges = {}
        self.lay = _layout(books)

    def grid(self, fid, sname):
        key = (fid, sname)
        if key not in self.grids:
            letters, rows, n, s = self.lay[key]
            cols = s.column_order()
            g = {}
            for r in s.sorted_rows():
                rn = rows[str(r.id)]
                for i, c in enumerate(cols):
                    if c in ("nguon", "raw_json"):
                        continue
                    g[(i + 1, rn)] = r.values.get(c)
            self.grids[key] = g
        return self.grids[key]

    def cell(self, fid, sname, col, row):
        v = self.grid(fid, sname).get((col, row))
        if isinstance(v, F):
            return self.value(fid, sname, v)
        return v

    def value(self, fid, sname, f: F):
        if f.value is not None or f.error is not None:
            if f.error:
                raise FormulaError(f.error)
            return f.value
        if f._busy:
            raise FormulaError("circular reference")
        f._busy = True
        try:
            toks = [m for m in TOKEN.finditer(f.text[1:]) if m.lastgroup != "ws"]
            p = _Parser(self, fid, sname, toks, f.text)
            v = p.expr()
            if p.i != len(toks):
                raise FormulaError(f"unparsed tail in {f.text}")
            if isinstance(v, Range):
                v = v[0] if len(v) == 1 else (_ for _ in ()).throw(FormulaError("range as a value"))
            f.value = v
            return v
        except FormulaError as e:
            f.error = str(e)
            raise
        except (ZeroDivisionError, ValueError, OverflowError) as e:
            f.error = f"{type(e).__name__}: {e}"
            raise FormulaError(f.error) from e
        finally:
            f._busy = False


def _num(v):
    # an empty cell is 0; an empty text ("" written as an inline string) is text, #VALUE! in arithmetic as in Excel
    if v is None:
        return 0.0
    if isinstance(v, bool):
        return 1.0 if v else 0.0
    if isinstance(v, (int, float)):
        return float(v)
    raise FormulaError(f"#VALUE! ({v!r} is not a number)")


def _bool(v):
    if isinstance(v, str):
        if v.upper() in ("TRUE", "FALSE"):
            return v.upper() == "TRUE"
        raise FormulaError(f"#VALUE! ({v!r} is not true / false)")
    return bool(_num(v))


def _cmp(a, b, op):
    if a is None:
        a = "" if isinstance(b, str) else 0
    if b is None:
        b = "" if isinstance(a, str) else 0
    if isinstance(a, str) or isinstance(b, str):
        if not (isinstance(a, str) and isinstance(b, str)):
            ka, kb = (1 if isinstance(a, str) else 0), (1 if isinstance(b, str) else 0)
            a, b = ka, kb
        else:
            a, b = a.lower(), b.lower()
    else:
        a, b = _num(a), _num(b)
    return {"=": a == b, "<>": a != b, "<": a < b, ">": a > b, "<=": a <= b, ">=": a >= b}[op]


def _numbers(args, scalars_count=True):
    out = []
    for a in args:
        if isinstance(a, Range):
            out.extend(float(x) for x in a if isinstance(x, (int, float)) and not isinstance(x, bool))
        elif isinstance(a, (int, float)) and not isinstance(a, bool):
            out.append(float(a))
        elif isinstance(a, bool) and scalars_count:
            out.append(1.0 if a else 0.0)
    return out


def _round_half_away(x, n=0):
    m = 10 ** int(n)
    y = abs(x) * m
    r = math.floor(y + 0.5 + 1e-12)
    return math.copysign(r / m, x)


def _median(xs):
    xs = sorted(xs)
    if not xs:
        raise FormulaError("#NUM! (MEDIAN of nothing)")
    k = len(xs) // 2
    return xs[k] if len(xs) % 2 else (xs[k - 1] + xs[k]) / 2


def _pairs(ys, xs):
    if not (isinstance(ys, Range) and isinstance(xs, Range)) or len(ys) != len(xs):
        raise FormulaError("#N/A (SLOPE needs two ranges of one size)")
    pts = [(float(x), float(y)) for x, y in zip(xs, ys)
           if isinstance(x, (int, float)) and isinstance(y, (int, float)) and not isinstance(x, bool) and not isinstance(y, bool)]
    if len(pts) < 2:
        raise FormulaError("#DIV/0! (fewer than two points)")
    mx = sum(p[0] for p in pts) / len(pts)
    my = sum(p[1] for p in pts) / len(pts)
    sxx = sum((p[0] - mx) ** 2 for p in pts)
    sxy = sum((p[0] - mx) * (p[1] - my) for p in pts)
    if sxx == 0:
        raise FormulaError("#DIV/0!")
    return sxy / sxx, mx, my


def _slope(ys, xs):
    return _pairs(ys, xs)[0]


def _intercept(ys, xs):
    b, mx, my = _pairs(ys, xs)
    return my - b * mx


def _ceiling(x, sig=1.0):
    x, sig = _num(x), _num(sig)
    if sig == 0:
        return 0.0
    return math.ceil(x / sig - 1e-12) * sig


def _floor(x, sig=1.0):
    x, sig = _num(x), _num(sig)
    if sig == 0:
        raise FormulaError("#DIV/0!")
    return math.floor(x / sig + 1e-12) * sig


def _index(rng, n, _col=None):
    n = int(_num(n))
    if not isinstance(rng, Range) or n < 1 or n > len(rng):
        raise FormulaError(f"#REF! (INDEX {n})")
    return rng[n - 1]


def _match(v, rng, kind=1):
    if int(_num(kind)) != 0:
        raise FormulaError("MATCH: only exact match (0) is supported")
    for i, x in enumerate(rng):
        if x is not None and _cmp(x, v, "="):
            return i + 1
    raise FormulaError(f"#N/A (MATCH {v!r})")


def _if(c, a=False, b=False):
    return a if _bool(c) else b


def _crit(c):
    """An Excel criterion (COUNTIFS / SUMIFS; MAXIFS is left out: Excel stores it as _xlfn.MAXIFS): a number matches equal numbers; a text with a leading <, >, <=, >=,
    = or <> compares; any other text matches case-insensitively with the wildcards * and ? (~ escapes)."""
    if isinstance(c, bool):
        return lambda x: isinstance(x, bool) and x == c
    if isinstance(c, (int, float)):
        return lambda x: isinstance(x, (int, float)) and not isinstance(x, bool) and float(x) == float(c)
    c = "" if c is None else str(c)
    m = re.match(r"^(<=|>=|<>|<|>|=)(.*)$", c, re.S)
    op, rest = (m.group(1), m.group(2)) if m else ("=", c)
    try:
        num = float(rest)
    except ValueError:
        num = None
    if op in ("<", ">", "<=", ">=") or (num is not None and m):
        if num is None:
            return lambda x: isinstance(x, str) and _cmp(x, rest, op)
        if op in ("=", "<>"):
            eq = lambda x: isinstance(x, (int, float)) and not isinstance(x, bool) and float(x) == num  # noqa: E731
            return eq if op == "=" else (lambda x: not eq(x))
        return lambda x: isinstance(x, (int, float)) and not isinstance(x, bool) and _cmp(float(x), num, op)
    rx = re.compile("^" + _wild(rest) + "$", re.I | re.S)
    if rest == "":
        hit = lambda x: x is None or x == ""  # noqa: E731
    else:
        hit = lambda x: x is not None and not isinstance(x, bool) and rx.match(_text(x)) is not None  # noqa: E731
    return hit if op == "=" else (lambda x: not hit(x))


def _wild(pattern: str) -> str:
    """An Excel wildcard pattern as a regex: * any run, ? one character, ~* ~? ~~ the character itself."""
    out, i = [], 0
    while i < len(pattern):
        ch = pattern[i]
        if ch == "~" and i + 1 < len(pattern) and pattern[i + 1] in "*?~":
            out.append(re.escape(pattern[i + 1]))
            i += 2
            continue
        out.append(".*" if ch == "*" else "." if ch == "?" else re.escape(ch))
        i += 1
    return "".join(out)


def _ifs_mask(pairs):
    if len(pairs) % 2:
        raise FormulaError("#VALUE! (criteria pairs)")
    rngs = pairs[0::2]
    if not all(isinstance(r, Range) for r in rngs) or len({len(r) for r in rngs}) != 1:
        raise FormulaError("#VALUE! (criteria ranges of one size)")
    tests = [_crit(c) for c in pairs[1::2]]
    return [all(t(r[i]) for r, t in zip(rngs, tests)) for i in range(len(rngs[0]))]


def _countifs(*pairs):
    return float(sum(_ifs_mask(list(pairs))))


def _agg_ifs(fn, empty):
    def run(values, *pairs):
        if not isinstance(values, Range):
            raise FormulaError("#VALUE! (a range of values)")
        mask = _ifs_mask(list(pairs))
        if len(mask) != len(values):
            raise FormulaError("#VALUE! (ranges of one size)")
        xs = [float(v) for v, ok in zip(values, mask) if ok and isinstance(v, (int, float)) and not isinstance(v, bool)]
        return fn(xs) if xs else empty
    return run


def _find(needle, hay, start=1):
    i = _text(hay).find(_text(needle), int(_num(start)) - 1)
    if i < 0:
        raise FormulaError("#VALUE! (FIND: not found)")
    return float(i + 1)


FUNCS = {
    "IF": _if,
    "AND": lambda *a: all(_bool(x) for x in a),
    "OR": lambda *a: any(_bool(x) for x in a),
    "NOT": lambda a: not _bool(a),
    "MAX": lambda *a: max(_numbers(a) or [0.0]),
    "MIN": lambda *a: min(_numbers(a) or [0.0]),
    "SUM": lambda *a: sum(_numbers(a)),
    "AVERAGE": lambda *a: (lambda xs: sum(xs) / len(xs) if xs else (_ for _ in ()).throw(FormulaError("#DIV/0!")))(_numbers(a)),
    "COUNT": lambda *a: float(len(_numbers(a, scalars_count=False))),
    "MEDIAN": lambda *a: _median(_numbers(a)),
    "SLOPE": _slope,
    "INTERCEPT": _intercept,
    "SQRT": lambda x: math.sqrt(_num(x)),
    "LN": lambda x: math.log(_num(x)),
    "EXP": lambda x: math.exp(_num(x)),
    "ABS": lambda x: abs(_num(x)),
    "POWER": lambda a, b: _num(a) ** _num(b),
    "ROUND": lambda x, n=0: _round_half_away(_num(x), _num(n)),
    "CEILING": _ceiling,
    "FLOOR": _floor,
    "INDEX": _index,
    "MATCH": _match,
    "ISNUMBER": lambda x: isinstance(x, (int, float)) and not isinstance(x, bool),
    "COUNTIFS": _countifs,
    "SUMIFS": _agg_ifs(sum, 0.0),
    "LEN": lambda x: float(len(_text(x))),
    "FIND": _find,
    "LEFT": lambda x, n=1: _text(x)[:max(0, int(_num(n)))],
    "SUBSTITUTE": lambda x, old, new: _text(x).replace(_text(old), _text(new)) if _text(old) else _text(x),
}
LAZY = {"IF", "IFERROR", "ISNUMBER"}


class _Parser:
    def __init__(self, ev: _Eval, fid, sname, toks, text):
        self.ev, self.fid, self.sname, self.toks, self.text, self.i = ev, fid, sname, toks, text, 0

    def peek(self):
        return self.toks[self.i] if self.i < len(self.toks) else None

    def take(self, kind=None, val=None):
        t = self.peek()
        if t is None or (kind and t.lastgroup != kind) or (val is not None and t.group() != val):
            raise FormulaError(f"syntax at token {self.i} of {self.text}")
        self.i += 1
        return t

    def is_op(self, *ops):
        t = self.peek()
        return t is not None and t.lastgroup == "op" and t.group() in ops

    def expr(self):
        a = self.concat()
        while self.is_op("=", "<>", "<", ">", "<=", ">="):
            op = self.take().group()
            a = _cmp(_scalar(a), _scalar(self.concat()), op)
        return a

    def concat(self):
        a = self.add()
        while self.is_op("&"):
            self.take()
            b = self.add()
            a = f"{_text(_scalar(a))}{_text(_scalar(b))}"
        return a

    def add(self):
        a = self.mul()
        while self.is_op("+", "-"):
            op = self.take().group()
            b = self.mul()
            a = _num(_scalar(a)) + _num(_scalar(b)) if op == "+" else _num(_scalar(a)) - _num(_scalar(b))
        return a

    def mul(self):
        a = self.pow()
        while self.is_op("*", "/"):
            op = self.take().group()
            b = _num(_scalar(self.pow()))
            if op == "*":
                a = _num(_scalar(a)) * b
            else:
                if b == 0:
                    raise FormulaError("#DIV/0!")
                a = _num(_scalar(a)) / b
        return a

    def pow(self):
        a = self.unary()
        while self.is_op("^"):
            self.take()
            a = _num(_scalar(a)) ** _num(_scalar(self.unary()))
        return a

    def unary(self):
        if self.is_op("-"):
            self.take()
            return -_num(_scalar(self.unary()))
        if self.is_op("+"):
            self.take()
            return self.unary()
        return self.primary()

    def primary(self):
        t = self.peek()
        if t is None:
            raise FormulaError(f"unexpected end of {self.text}")
        g = t.lastgroup
        if g == "num":
            self.take()
            return float(t.group())
        if g == "str":
            self.take()
            return t.group()[1:-1].replace('""', '"')
        if g == "bool":
            self.take()
            return t.group() == "TRUE"
        if g == "ref":
            self.take()
            return self.reference(t.group())
        if g == "func":
            name = self.take().group()
            self.take("op", "(")
            args_src = []
            if not self.is_op(")"):
                while True:
                    args_src.append(self.i)
                    self.skip_arg()
                    if self.is_op(","):
                        self.take()
                        continue
                    break
            end = self.i
            self.take("op", ")")
            fn = FUNCS.get(name)
            if fn is None:
                raise FormulaError(f"function {name} is not in the evaluator's subset")
            if name in ("IFERROR", "ISNUMBER"):
                # an error inside is caught (Excel: IFERROR gives its second argument, ISNUMBER FALSE)
                try:
                    v = self.sub(args_src[0], args_src[1], True) if len(args_src) > 1 else self.sub(args_src[0], end)
                    v = _scalar(v)
                except (FormulaError, ZeroDivisionError, ValueError, OverflowError):
                    if name == "ISNUMBER":
                        return False
                    return self.sub(args_src[1], end) if len(args_src) > 1 else 0.0
                return fn(v) if name == "ISNUMBER" else v
            if name in LAZY:
                c = self.sub(args_src[0], args_src[1], True) if len(args_src) > 1 else self.sub(args_src[0], end)
                pick = 1 if _bool(_scalar(c)) else 2
                if pick < len(args_src):
                    stop = args_src[pick + 1] if pick + 1 < len(args_src) else end
                    return self.sub(args_src[pick], stop, trailing_comma=pick + 1 < len(args_src))
                return False
            vals = []
            for k, start in enumerate(args_src):
                stop = args_src[k + 1] if k + 1 < len(args_src) else end
                vals.append(self.sub(start, stop, trailing_comma=k + 1 < len(args_src)))
            return fn(*vals)
        if g == "op" and t.group() == "(":
            self.take()
            v = self.expr()
            self.take("op", ")")
            return v
        raise FormulaError(f"unexpected {t.group()!r} in {self.text}")

    def skip_arg(self):
        depth = 0
        while self.i < len(self.toks):
            t = self.toks[self.i]
            if t.lastgroup == "op" and t.group() == "(":
                depth += 1
            elif t.lastgroup == "op" and t.group() == ")":
                if depth == 0:
                    return
                depth -= 1
            elif t.lastgroup == "op" and t.group() == "," and depth == 0:
                return
            elif t.lastgroup == "func":
                pass
            self.i += 1

    def sub(self, start, stop, trailing_comma=False):
        stop = stop - 1 if trailing_comma else stop
        p = _Parser(self.ev, self.fid, self.sname, self.toks[start:stop], self.text)
        v = p.expr()
        if p.i != len(p.toks):
            raise FormulaError(f"argument not fully read in {self.text}")
        return v

    def reference(self, text):
        sheet = self.sname
        if "!" in text:
            sheet, text = text.split("!", 1)
        parts = text.split(":")
        cells = [CELL.fullmatch(p) for p in parts]
        if any(c is None for c in cells):
            raise FormulaError(f"bad reference {text}")
        if (self.fid, sheet) not in self.ev.lay:
            raise FormulaError(f"#REF! no sheet {sheet}")
        if len(cells) == 1:
            return self.ev.cell(self.fid, sheet, _col_num(cells[0].group(1)), int(cells[0].group(2)))
        c1, r1 = _col_num(cells[0].group(1)), int(cells[0].group(2))
        c2, r2 = _col_num(cells[1].group(1)), int(cells[1].group(2))
        key = (self.fid, sheet, c1, r1, c2, r2)
        if key in self.ev.ranges:
            return self.ev.ranges[key]
        out = Range()
        for c in range(c1, c2 + 1):
            for r in range(r1, r2 + 1):
                out.append(self.ev.cell(self.fid, sheet, c, r))
        self.ev.ranges[key] = out
        return out


def _scalar(v):
    if isinstance(v, Range):
        if len(v) == 1:
            return v[0]
        raise FormulaError("#VALUE! (a range where one value is needed)")
    return v


def _text(v):
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return "" if v is None else str(v)


def evaluate(books) -> None:
    ev = _Eval(books)
    for fid, book in books.items():
        for sname, s in book.sheets.items():
            for r in s.rows.values():
                for v in r.values.values():
                    if isinstance(v, F) and v.text and v.value is None and v.error is None:
                        try:
                            ev.value(fid, sname, v)
                        except FormulaError:
                            pass


# ---------------------------------------------------------------------------------------------------------- tests
def _same(a, b) -> tuple[bool, float]:
    if isinstance(b, bool) or isinstance(a, bool):
        if isinstance(a, (bool, int, float)) and isinstance(b, (bool, int, float)):
            return bool(a) == bool(b), 0.0
        return False, 0.0
    if isinstance(a, (int, float)) and isinstance(b, (int, float)):
        d = abs(float(a) - float(b))
        return d <= TOL * max(1.0, abs(float(b))), d
    if a is None:
        a = ""
    if b is None:
        b = ""
    return str(a).lower() == str(b).lower(), 0.0


def check(books, copies: list[dict]) -> list[dict]:
    """One result a formula column (formula == expect on every row) and a copied input column (copy == source)."""
    out = []
    for fid in sorted(books):
        for sname, s in books[fid].sheets.items():
            per = {}
            for r in s.sorted_rows():
                for cname, v in r.values.items():
                    if not isinstance(v, F):
                        continue
                    st = per.setdefault(cname, {"n": 0, "ok": 0, "errors": 0, "max_diff": 0.0, "examples": [], "refs": set(),
                                                "unchecked": 0})
                    st["n"] += 1
                    if v.ref:
                        st["refs"].add(v.ref)
                    if v.error:
                        st["errors"] += 1
                        if len(st["examples"]) < 5:
                            st["examples"].append(f"{r.id}: {v.error}")
                        continue
                    if v.expect is None:
                        st["unchecked"] += 1
                        continue
                    ok, d = _same(v.value, v.expect)
                    st["max_diff"] = max(st["max_diff"], d)
                    if ok:
                        st["ok"] += 1
                    elif len(st["examples"]) < 5:
                        st["examples"].append(f"{r.id}: {v.value!r} != {v.expect!r}")
            for cname, st in sorted(per.items()):
                checked = st["n"] - st["unchecked"]
                status = "OK" if st["ok"] == checked and not st["errors"] else "FAIL"
                out.append({"kind": "cong_thuc", "file": fid, "sheet": sname, "column": cname, "rows": st["n"],
                            "match": st["ok"], "unchecked": st["unchecked"], "errors": st["errors"],
                            "max_diff": st["max_diff"], "reference": "; ".join(sorted(st["refs"])), "status": status,
                            "examples": "; ".join(st["examples"])})
    for cp in copies:
        book, src_book = books.get(cp["file"]), books.get(cp["src_file"])
        s = book.sheets[cp["sheet"]] if book else None
        for col, src_col in cp["cols"].items():
            res = {"kind": "input", "file": cp["file"], "sheet": cp["sheet"], "column": col, "rows": 0, "match": 0,
                   "unchecked": 0, "errors": 0, "max_diff": 0.0,
                   "reference": f"{cp['src_file']}/{cp['src_sheet']}.{src_col}", "status": "OK", "examples": ""}
            if src_book is None or cp["src_sheet"] not in src_book.sheets:
                res["status"] = "PENDING"
                res["examples"] = "file nguồn chưa xuất"
                out.append(res)
                continue
            src = src_book.sheets[cp["src_sheet"]]
            bad = []
            for r in s.sorted_rows():
                res["rows"] += 1
                sr = src.rows.get(r.id)
                if sr is None:
                    bad.append(f"{r.id}: không có ở nguồn")
                    continue
                a, b = r.values.get(col), sr.values.get(src_col)
                a = a.value if isinstance(a, F) else a
                b = b.value if isinstance(b, F) else b
                ok, d = _same(a, b)
                res["max_diff"] = max(res["max_diff"], d)
                if ok:
                    res["match"] += 1
                else:
                    bad.append(f"{r.id}: {a!r} != {b!r}")
            if bad:
                res["status"] = "FAIL"
                res["errors"] = len(bad)
                res["examples"] = "; ".join(bad[:5])
            out.append(res)
    return out


def run(ctx) -> list[dict]:
    """Resolve, evaluate and test every formula and copy of ctx.books; the results also land on ctx.formula_checks."""
    problems = resolve(ctx.books)
    for p in problems[:50]:
        ctx.issue(f"formula: {p}")
    evaluate(ctx.books)
    results = check(ctx.books, getattr(ctx, "copies", []))
    ctx.formula_checks = results
    return results
