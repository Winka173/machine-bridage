"""Writers: xlsx (openpyxl, fixed metadata and zip dates so equal data gives equal bytes) and CSV (UTF-8, comma, header)."""
from __future__ import annotations

import csv
import io
import re
import zipfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

FIXED_DATE = "2000-01-01T00:00:00Z"
ZIP_DATE = (1980, 1, 1, 0, 0, 0)
HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="2F4F4F")
MAX_WIDTH = 60


def _cell_value(v):
    if isinstance(v, str):
        return ILLEGAL_CHARACTERS_RE.sub("", v)
    return v


def write_xlsx(path: Path, sheets: list[tuple[str, list[str], list[list], set]]):
    """sheets: [(name, header, rows, formula columns)]. A text starting with '=' is stored as text unless its column is a
    formula column (layer B)."""
    wb = Workbook(write_only=True)
    for name, header, rows, formula_cols in sheets:
        ws = wb.create_sheet(title=name)
        widths = [min(MAX_WIDTH, max(8, len(h) + 2)) for h in header]
        for line in rows[:200]:
            for i, v in enumerate(line):
                if v is not None:
                    widths[i] = min(MAX_WIDTH, max(widths[i], len(str(v)) + 1))
        for i, w in enumerate(widths):
            ws.column_dimensions[get_column_letter(i + 1)].width = w
        ws.freeze_panes = "B2"
        if header:
            ws.auto_filter.ref = f"A1:{get_column_letter(len(header))}{len(rows) + 1}"
        head = []
        for h in header:
            c = WriteOnlyCell(ws, value=h)
            c.font = HEADER_FONT
            c.fill = HEADER_FILL
            c.alignment = Alignment(vertical="top")
            head.append(c)
        ws.append(head)
        fcols = {i for i, h in enumerate(header) if h in formula_cols}
        for line in rows:
            out = []
            for i, v in enumerate(line):
                v = _cell_value(v)
                if isinstance(v, str) and v.startswith("=") and i not in fcols:
                    c = WriteOnlyCell(ws, value=v)
                    c.data_type = "s"
                    out.append(c)
                else:
                    out.append(v)
            ws.append(out)
    wb.properties.creator = "Tools/export"
    wb.calculation.fullCalcOnLoad = True  # layer B formulas carry no cached value: Excel / LibreOffice compute them on open
    buf = io.BytesIO()
    wb.save(buf)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(normalise_zip(buf.getvalue()))


def normalise_zip(data: bytes) -> bytes:
    """The same entries with a fixed date and fixed core properties (created / modified), in the same order."""
    src = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            body = src.read(info)
            if info.filename == "docProps/core.xml":
                text = body.decode("utf-8")
                text = re.sub(r"(<dcterms:created[^>]*>)[^<]*(</dcterms:created>)", rf"\g<1>{FIXED_DATE}\g<2>", text)
                text = re.sub(r"(<dcterms:modified[^>]*>)[^<]*(</dcterms:modified>)", rf"\g<1>{FIXED_DATE}\g<2>", text)
                text = re.sub(r"<cp:lastModifiedBy>[^<]*</cp:lastModifiedBy>", "", text)
                body = text.encode("utf-8")
            zi = zipfile.ZipInfo(info.filename, date_time=ZIP_DATE)
            zi.compress_type = zipfile.ZIP_DEFLATED
            zi.external_attr = 0o600 << 16
            zi.create_system = 0
            dst.writestr(zi, body)
    return out.getvalue()


def csv_value(v) -> str:
    if v is None:
        return ""
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, float):
        return repr(v)
    return str(v)


def write_csv(path: Path, header: list[str], rows: list[list]):
    path.parent.mkdir(parents=True, exist_ok=True)
    buf = io.StringIO()
    w = csv.writer(buf, lineterminator="\n")
    w.writerow(header)
    for line in rows:
        w.writerow([csv_value(v) for v in line])
    path.write_bytes(buf.getvalue().encode("utf-8"))
